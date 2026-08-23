import os
import sys
import json
import re
import site

# 1. Automatically ensure User Site Packages are loaded even if Apache/PHP runs without APPDATA
user_site = site.getusersitepackages()
if user_site and os.path.exists(user_site) and user_site not in sys.path:
    sys.path.insert(0, user_site)

# Also explicitly look for Windows standard AppData Python directories
for user_base in [os.environ.get('USERPROFILE', ''), r'C:\Users\Yogi Nugraha', r'C:\Users\Default']:
    if user_base:
        for py_ver in ['Python313', 'Python312', 'Python311', 'Python310']:
            cand = os.path.join(user_base, 'AppData', 'Roaming', 'Python', py_ver, 'site-packages')
            if os.path.exists(cand) and cand not in sys.path:
                sys.path.insert(0, cand)

def parse_cell_numeric(val):
    if not val:
        return ""
    s_val = str(val).strip()
    clean_num = s_val.replace(',', '').replace(' ', '')
    if re.match(r'^-?\d+(\.\d+)?$', clean_num):
        try:
            return float(clean_num) if '.' in clean_num else int(clean_num)
        except ValueError:
            return s_val
    return s_val

def convert_with_pdfplumber(input_pdf_path, output_xlsx_path):
    import pdfplumber
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    total_tables_found = 0
    header_font = Font(name="Calibri", size=11, bold=True, color="1F2937")
    header_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    cell_font = Font(name="Calibri", size=11, color="111827")
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    with pdfplumber.open(input_pdf_path) as pdf:
        total_pages = len(pdf.pages)
        if total_pages == 0:
            raise ValueError("Dokumen PDF tidak memiliki halaman.")

        for page_idx, page in enumerate(pdf.pages, start=1):
            sheet_title = f"Page {page_idx}" if total_pages > 1 else "Sheet1"
            ws = wb.create_sheet(title=sheet_title)
            ws.views.sheetView[0].showGridLines = True

            tables = page.extract_tables({
                "vertical_strategy": "lines",
                "horizontal_strategy": "lines",
                "snap_tolerance": 4,
                "join_tolerance": 4,
            })

            if not tables:
                tables = page.extract_tables({
                    "vertical_strategy": "text",
                    "horizontal_strategy": "text",
                    "snap_tolerance": 4,
                    "join_tolerance": 4,
                    "text_x_tolerance": 3,
                    "text_y_tolerance": 3,
                })

            current_row = 1

            if tables:
                for table in tables:
                    if not table:
                        continue
                    total_tables_found += 1
                    is_first_row = True
                    for row_data in table:
                        if not row_data or all(c is None or str(c).strip() == '' for c in row_data):
                            continue
                        
                        for col_idx, cell_value in enumerate(row_data, start=1):
                            val_to_write = parse_cell_numeric(cell_value)
                            cell = ws.cell(row=current_row, column=col_idx, value=val_to_write)
                            cell.font = header_font if is_first_row else cell_font
                            cell.border = thin_border
                            if is_first_row:
                                cell.fill = header_fill
                                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                            else:
                                cell.alignment = Alignment(
                                    horizontal="right" if isinstance(val_to_write, (int, float)) else "left",
                                    vertical="center",
                                    wrap_text=True
                                )

                        current_row += 1
                        is_first_row = False
                    current_row += 1
            else:
                text = page.extract_text(layout=True) or ""
                lines = text.splitlines()
                for line in lines:
                    if not line.strip():
                        continue
                    cols = [c.strip() for c in re.split(r'\s{2,}|\t', line) if c.strip()]
                    if not cols:
                        cols = [line.strip()]

                    for col_idx, val in enumerate(cols, start=1):
                        val_to_write = parse_cell_numeric(val)
                        cell = ws.cell(row=current_row, column=col_idx, value=val_to_write)
                        cell.font = cell_font
                        cell.alignment = Alignment(
                            horizontal="right" if isinstance(val_to_write, (int, float)) else "left",
                            vertical="center"
                        )
                    current_row += 1

            # Auto fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        for l in str(cell.value).split('\n'):
                            max_len = max(max_len, len(l))
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)

    if len(wb.sheetnames) == 0:
        ws = wb.create_sheet(title="Sheet1")
        ws.cell(row=1, column=1, value="No table data detected in PDF")

    wb.save(output_xlsx_path)
    return total_pages, total_tables_found

def convert_with_pymupdf_xlsxwriter(input_pdf_path, output_xlsx_path):
    import fitz
    import xlsxwriter

    doc = fitz.open(input_pdf_path)
    total_pages = len(doc)
    total_tables_found = 0

    wb = xlsxwriter.Workbook(output_xlsx_path)
    header_format = wb.add_format({
        'bold': True,
        'bg_color': '#F3F4F6',
        'border': 1,
        'border_color': '#E5E7EB',
        'align': 'center',
        'valign': 'vcenter'
    })
    cell_format_left = wb.add_format({
        'border': 1,
        'border_color': '#E5E7EB',
        'align': 'left',
        'valign': 'vcenter'
    })
    cell_format_right = wb.add_format({
        'border': 1,
        'border_color': '#E5E7EB',
        'align': 'right',
        'valign': 'vcenter'
    })

    for page_idx in range(total_pages):
        page = doc[page_idx]
        sheet_title = f"Page {page_idx + 1}" if total_pages > 1 else "Sheet1"
        ws = wb.add_worksheet(sheet_title)
        ws.show_gridlines(True)

        current_row = 0
        tables_found_on_page = False

        if hasattr(page, 'find_tables'):
            try:
                table_finder = page.find_tables()
                if table_finder and table_finder.tables:
                    tables_found_on_page = True
                    for table in table_finder.tables:
                        extracted = table.extract()
                        if not extracted:
                            continue
                        total_tables_found += 1
                        for r_idx, row in enumerate(extracted):
                            for c_idx, cell in enumerate(row):
                                val = parse_cell_numeric(cell)
                                fmt = header_format if r_idx == 0 else (cell_format_right if isinstance(val, (int, float)) else cell_format_left)
                                ws.write(current_row, c_idx, val, fmt)
                            current_row += 1
                        current_row += 1
            except Exception:
                tables_found_on_page = False

        if not tables_found_on_page:
            text = page.get_text("text") or ""
            for line in text.splitlines():
                if not line.strip():
                    continue
                cols = [c.strip() for c in re.split(r'\s{2,}|\t', line) if c.strip()]
                if not cols:
                    cols = [line.strip()]
                for c_idx, val in enumerate(cols):
                    p_val = parse_cell_numeric(val)
                    fmt = cell_format_right if isinstance(p_val, (int, float)) else cell_format_left
                    ws.write(current_row, c_idx, p_val, fmt)
                current_row += 1

    wb.close()
    doc.close()
    return total_pages, total_tables_found

def convert_pdf_to_xlsx(input_pdf_path, output_xlsx_path):
    if not os.path.exists(input_pdf_path):
        return {"status": "error", "message": f"File PDF tidak ditemukan: {input_pdf_path}"}

    # Primary engine: pdfplumber + openpyxl
    try:
        total_pages, total_tables = convert_with_pdfplumber(input_pdf_path, output_xlsx_path)
        return {
            "status": "success",
            "page_count": total_pages,
            "table_count": total_tables,
            "engine": "pdfplumber",
            "output_size": os.path.getsize(output_xlsx_path)
        }
    except Exception as e1:
        # Secondary engine: pymupdf + xlsxwriter fallback
        try:
            total_pages, total_tables = convert_with_pymupdf_xlsxwriter(input_pdf_path, output_xlsx_path)
            return {
                "status": "success",
                "page_count": total_pages,
                "table_count": total_tables,
                "engine": "pymupdf_fallback",
                "output_size": os.path.getsize(output_xlsx_path)
            }
        except Exception as e2:
            return {"status": "error", "message": f"Gagal mengekstrak PDF ke Excel: {str(e1)} (Fallback error: {str(e2)})"}

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(json.dumps({"status": "error", "message": "Usage: python pdf2xlsx.py <input_pdf> <output_xlsx>"}))
        sys.exit(1)

    input_pdf = os.path.abspath(sys.argv[1])
    output_xlsx = os.path.abspath(sys.argv[2])

    result = convert_pdf_to_xlsx(input_pdf, output_xlsx)
    print(json.dumps(result))
    
    if result.get("status") == "success":
        sys.exit(0)
    else:
        sys.exit(1)
